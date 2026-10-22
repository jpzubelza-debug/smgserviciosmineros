"""Pruebas de la importación automática de Clientes.xlsx al arrancar la aplicación.

Cada "arranque" es un proceso nuevo que importa main (lo mismo que hace uvicorn --reload), con una carpeta de
aplicación temporal (main.py, schema.sql, certificado_pdf.py y, según el caso, Clientes.xlsx) y DATA_DIR temporal.
La base con clientes existentes es una COPIA de dashboard.db (la original se abre en solo lectura; se elige con
SMG_TEST_DB_ORIGEN). Ninguna prueba escribe en la base real.

Uso: python -m unittest tests.test_clientes_import -v
"""
import contextlib
import hashlib
import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import unittest

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_REAL = os.path.abspath(os.environ.get("SMG_TEST_DB_ORIGEN") or os.path.join(BASE_DIR, "dashboard.db"))
EXCEL = os.path.join(BASE_DIR, "Clientes.xlsx")
TMP = tempfile.mkdtemp(prefix="smg_test_clientes_")


def _clientes_en_excel():
    from openpyxl import load_workbook
    libro = load_workbook(EXCEL, read_only=True, data_only=True)
    hoja = libro["Clientes"] if "Clientes" in libro.sheetnames else libro.active
    filas = hoja.iter_rows(values_only=True)
    next(filas, None)
    total = sum(1 for fila in filas if str((list(fila) + [None])[0] or "").strip())
    libro.close()
    return total


def _carpeta(nombre, excel=True, contenido_excel=None):
    """Carpeta de aplicación temporal y su DATA_DIR."""
    app = os.path.join(TMP, nombre, "app")
    datos = os.path.join(TMP, nombre, "datos")
    os.makedirs(os.path.join(app, "Imagenes"), exist_ok=True)
    os.makedirs(datos, exist_ok=True)
    for archivo in ("main.py", "schema.sql", "certificado_pdf.py", "verificacion_certificados.py", "facturacion.py", "facturacion_pdf.py"):
        shutil.copy(os.path.join(BASE_DIR, archivo), app)
    if contenido_excel is not None:
        with open(os.path.join(app, "Clientes.xlsx"), "wb") as f:
            f.write(contenido_excel)
    elif excel:
        shutil.copy(EXCEL, app)
    return app, datos


def _entorno(datos):
    entorno = {k: v for k, v in os.environ.items() if k not in ("DATA_DIR", "OPENAI_API_KEY")}
    entorno.update({"DATA_DIR": datos, "PYTHONIOENCODING": "utf-8"})
    return entorno


def _arrancar(app, datos, codigo="import main"):
    r = subprocess.run([sys.executable, "-c", codigo], cwd=app, env=_entorno(datos), capture_output=True,
                       text=True, encoding="utf-8", timeout=300)
    if r.returncode != 0:
        raise AssertionError(f"El arranque falló ({r.returncode}):\n{r.stdout}\n{r.stderr}")
    return r.stdout


@contextlib.contextmanager
def _db(datos):
    conn = sqlite3.connect(os.path.join(datos, "dashboard.db"), timeout=20)
    conn.row_factory = sqlite3.Row
    try:
        with conn:
            yield conn
    finally:
        conn.close()


def _clientes(datos):
    with _db(datos) as conn:
        return [tuple(f) for f in conn.execute("SELECT * FROM clientes ORDER BY codigo")]


def _huella(datos):
    """Contenido de todas las tablas y el esquema, para verificar que un arranque no cambia nada."""
    conn = sqlite3.connect(f"file:{os.path.join(datos, 'dashboard.db')}?mode=ro", uri=True)
    huella = {"esquema": sorted(map(tuple, conn.execute("SELECT type, name, sql FROM sqlite_master")))}
    for (tabla,) in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table' AND name NOT LIKE 'sqlite_%'"):
        filas = conn.execute(f'SELECT * FROM "{tabla}" ORDER BY rowid').fetchall()
        huella[tabla] = (len(filas), hashlib.sha256(repr(filas).encode()).hexdigest())
    conn.close()
    return huella


def _tabla_vacia(nombre, excel=True, contenido_excel=None):
    """Base con el esquema completo y la tabla clientes existente pero vacía."""
    app, datos = _carpeta(nombre, excel=False)
    salida = _arrancar(app, datos)
    assert "la tabla clientes esta vacia y no se encontro" in salida, salida
    if contenido_excel is not None:
        with open(os.path.join(app, "Clientes.xlsx"), "wb") as f:
            f.write(contenido_excel)
    elif excel:
        shutil.copy(EXCEL, app)
    return app, datos


TOTAL = _clientes_en_excel()


class ImportacionClientes(unittest.TestCase):
    def test_1_base_nueva_importa(self):
        app, datos = _carpeta("nueva")
        self.assertFalse(os.path.exists(os.path.join(datos, "dashboard.db")))
        salida = _arrancar(app, datos)
        self.assertIn(f"importacion inicial ejecutada: {TOTAL} clientes", salida)
        self.assertEqual(len(_clientes(datos)), TOTAL)

    def test_2_tabla_vacia_importa(self):
        app, datos = _tabla_vacia("vacia")
        self.assertEqual(_clientes(datos), [])
        salida = _arrancar(app, datos)
        self.assertIn(f"importacion inicial ejecutada: {TOTAL} clientes (la tabla clientes estaba vacia)", salida)
        self.assertEqual(len(_clientes(datos)), TOTAL)

    def test_3_existentes_no_se_modifican_en_reinicios(self):
        app, datos = _carpeta("existentes")
        origen = sqlite3.connect(f"file:{DB_REAL}?mode=ro", uri=True)
        destino = sqlite3.connect(os.path.join(datos, "dashboard.db"))
        origen.backup(destino)
        destino.close()
        origen.close()
        # Un cliente editado en la aplicación y con fecha antigua: el código anterior lo pisaba con el Excel en cada arranque.
        with _db(datos) as conn:
            codigo = conn.execute("SELECT codigo FROM clientes ORDER BY codigo LIMIT 1").fetchone()[0]
            conn.execute("UPDATE clientes SET razon_social = 'EDITADO EN LA APP', updated_at = '2000-01-01 00:00:00' WHERE codigo = ?",
                         (codigo,))
        _arrancar(app, datos)  # primer arranque: aplica las migraciones idempotentes que la copia aún necesite
        clientes, huella = _clientes(datos), _huella(datos)
        self.assertGreater(len(clientes), 0)
        for _ in range(3):
            salida = _arrancar(app, datos)
            self.assertIn(f"importacion automatica omitida: la tabla clientes ya tiene {len(clientes)} registros", salida)
            self.assertNotIn("Clientes migrados desde Excel", salida)
            self.assertEqual(_clientes(datos), clientes)  # incluye updated_at
            self.assertEqual(_huella(datos), huella)  # ni esquema ni datos de ninguna tabla cambian
        with _db(datos) as conn:
            fila = conn.execute("SELECT razon_social, updated_at FROM clientes WHERE codigo = ?", (codigo,)).fetchone()
        self.assertEqual(tuple(fila), ("EDITADO EN LA APP", "2000-01-01 00:00:00"))

    def test_4_importacion_deliberada_sigue_funcionando(self):
        app, datos = _carpeta("manual")
        _arrancar(app, datos)
        with _db(datos) as conn:
            codigo, original = conn.execute("SELECT codigo, razon_social FROM clientes ORDER BY codigo LIMIT 1").fetchone()
            conn.execute("UPDATE clientes SET razon_social = 'DESACTUALIZADO', updated_at = '2000-01-01 00:00:00' WHERE codigo = ?",
                         (codigo,))
        salida = _arrancar(app, datos, "import main\nconn = main.get_sqlite_connection()\n"
                                       "print('importados:', main.migrar_clientes_excel(conn))\nconn.commit()")
        self.assertIn(f"importados: {TOTAL}", salida)
        with _db(datos) as conn:
            fila = conn.execute("SELECT razon_social, updated_at FROM clientes WHERE codigo = ?", (codigo,)).fetchone()
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM clientes").fetchone()[0], TOTAL)
        self.assertEqual(fila["razon_social"], original)
        self.assertNotEqual(fila["updated_at"], "2000-01-01 00:00:00")

    def test_5_excel_ausente_mensaje_claro(self):
        app, datos = _carpeta("ausente", excel=False)
        salida = _arrancar(app, datos)
        self.assertIn("importacion inicial omitida: la tabla clientes esta vacia y no se encontro", salida)
        self.assertEqual(_clientes(datos), [])
        # El resto de la inicialización se completó.
        with _db(datos) as conn:
            self.assertTrue(conn.execute("SELECT 1 FROM sqlite_master WHERE name = 'reajuste_formulas'").fetchone())

    def test_6_excel_ilegible_se_revierte(self):
        app, datos = _tabla_vacia("ilegible", contenido_excel=b"esto no es un libro de Excel")
        salida = _arrancar(app, datos)
        self.assertIn("Error migrando Clientes.xlsx", salida)
        self.assertIn("importacion inicial fallida: se revirtio y la tabla clientes sigue vacia", salida)
        self.assertEqual(_clientes(datos), [])

    def test_7_falla_a_mitad_no_deja_carga_parcial(self):
        app, datos = _tabla_vacia("parcial", excel=False)
        # Se inserta un cliente y luego la lectura falla: la carga inicial debe revertirse completa.
        codigo = (
            "import main\n"
            f"main.CLIENTES_EXCEL_PATH = {EXCEL!r}\n"
            "def falla(conn):\n"
            "    conn.execute(\"INSERT INTO clientes (codigo, razon_social) VALUES ('X-PARCIAL', 'Parcial')\")\n"
            "    print('Error migrando Clientes.xlsx: simulado')\n"
            "    return None\n"
            "main.migrar_clientes_excel = falla\n"
            "conn = main.get_sqlite_connection()\n"
            "main.importar_clientes_inicial(conn)\n"
        )
        salida = _arrancar(app, datos, codigo)
        self.assertIn("importacion inicial fallida: se revirtio", salida)
        self.assertEqual(_clientes(datos), [])

    def test_8_arranques_simultaneos_importan_una_sola_vez(self):
        for ronda in range(2):
            app, datos = _tabla_vacia(f"simultaneos_{ronda}")
            procesos = [subprocess.Popen([sys.executable, "-c", "import main"], cwd=app, env=_entorno(datos),
                                         stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8")
                        for _ in range(3)]
            salidas = []
            for p in procesos:
                salida, error = p.communicate(timeout=300)
                self.assertEqual(p.returncode, 0, error)
                salidas.append(salida)
            ejecutadas = sum("importacion inicial ejecutada" in s for s in salidas)
            omitidas = sum("importacion automatica omitida" in s for s in salidas)
            self.assertEqual((ejecutadas, omitidas), (1, 2), salidas)
            self.assertEqual(len(_clientes(datos)), TOTAL)


if __name__ == "__main__":
    unittest.main()
