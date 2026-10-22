"""Prueba de punta a punta del pegado del resumen de cuenta de Tango (navegador real, base copiada).

Uso: python tests/e2e_tango.py [carpeta_capturas]
  SMG_TANGO_XLSX  resumen de cuenta exportado de Tango (por defecto, el de GE.MAR en la carpeta del proyecto).
  SMG_TANGO_TSV   opcional: texto copiado de ese Excel con el portapapeles real, para probarlo tal cual.
Lee el resumen, arma el texto que Excel deja en el portapapeles (formato contable y fechas día/mes o mes/día) y verifica:
el reconocimiento de cada fila contra el Excel, los errores de formato, que pegar y preparar formularios no guarda nada,
y que cada registro se confirma en su formulario con las validaciones de siempre.
La fecha de inicio 2026-02-28 es SOLO un dato de prueba.
"""
import datetime
import hashlib
import json
import os
import socket
import sqlite3
import sys
import tempfile
import threading
import time

sys.stdout.reconfigure(encoding="utf-8")  # consola de Windows (cp1252)
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)
CAPTURAS = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else os.path.join(tempfile.gettempdir(), "smg_e2e_tango_capturas"))
os.makedirs(CAPTURAS, exist_ok=True)
XLSX = os.path.abspath(os.environ.get("SMG_TANGO_XLSX") or os.path.join(BASE_DIR, "20260717 - Resumen de Cuenta GeMar 2025-2026 (1).xlsx"))
TSV_REAL = os.environ.get("SMG_TANGO_TSV")
if not os.path.exists(XLSX):
    sys.exit(f"No se encontró el resumen de Tango: {XLSX} (indicalo con SMG_TANGO_XLSX)")
DB_ORIGEN = os.path.abspath(os.environ.get("SMG_TEST_DB_ORIGEN") or os.path.join(BASE_DIR, "dashboard.db"))
SHA_ORIGEN = hashlib.sha256(open(DB_ORIGEN, "rb").read()).hexdigest()
SHA_XLSX = hashlib.sha256(open(XLSX, "rb").read()).hexdigest()
TMP = tempfile.mkdtemp(prefix="smg_e2e_tango_")
DB = os.path.join(TMP, "dashboard.db")
origen = sqlite3.connect(f"file:{DB_ORIGEN}?mode=ro", uri=True)
destino = sqlite3.connect(DB)
origen.backup(destino)
destino.close()
origen.close()
os.environ.update({"DATA_DIR": TMP, "APP_LOGIN_USER": "contingencia_e2e", "APP_LOGIN_PASSWORD": "clave-contingencia-e2e"})

import main  # noqa: E402
import openpyxl  # noqa: E402
import uvicorn  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402

assert os.path.abspath(main.SQLITE_DB_PATH) == os.path.abspath(DB), "La prueba no apunta a la copia"
sock = socket.socket()
sock.bind(("127.0.0.1", 0))
PUERTO = sock.getsockname()[1]
sock.close()
servidor = uvicorn.Server(uvicorn.Config(main.app, host="127.0.0.1", port=PUERTO, log_level="warning", lifespan="off"))
threading.Thread(target=servidor.run, daemon=True).start()
while not servidor.started:
    time.sleep(0.05)
URL = f"http://127.0.0.1:{PUERTO}"
API = "/api/facturacion"
CLAVE = "Clave-Prueba-1"
CORTE = "2026-02-28"


def db():
    conn = sqlite3.connect(DB, timeout=20)
    conn.row_factory = sqlite3.Row
    return conn


# ---------------------------------------------------------------- el Excel: valores esperados y texto del portapapeles
HOJA = openpyxl.load_workbook(XLSX, data_only=True).worksheets[0]
ENCABEZADOS = [c.value for c in HOJA[1]]
TIPOS = {"FAC": "FC", "N/C": "NC", "N/D": "ND", "REC": "REC"}


def iso(v):
    return v.strftime("%Y-%m-%d") if isinstance(v, datetime.datetime) else None


ESPERADO = []
for fila in HOJA.iter_rows(min_row=2, values_only=True):
    x = dict(zip(ENCABEZADOS, fila))
    nc = str(x["N_COMP"]).strip()
    ESPERADO.append({"tipo": TIPOS[x["T_COMP"]], "letra": nc[0] if nc[0].isalpha() else None, "punto_venta": int(nc.lstrip("ABCEM")[:5]),
                     "numero": int(nc[-8:]), "fecha": iso(x["FECHA_EMIS"]), "vencimiento": iso(x["FECHA_VTO"]),
                     "debe": round(x["DEBE"], 2), "haber": round(x["HABER"], 2), "saldo": round(x["SALDO"], 2), "estado": x["ESTADO"],
                     "leyenda": (x["LEYENDA"] or "").strip()})
RAZON = HOJA["A2"].value


def portapapeles(regional="es-AR"):
    """Texto que Excel deja en el portapapeles al copiar la hoja: formato contable y fechas según la configuración regional.
    Para es-AR coincide byte a byte con una copia real hecha con Excel."""
    lineas = []
    for fila in HOJA.iter_rows():
        celdas = []
        for c in fila:
            v, contable = c.value, "#,##0.00" in (c.number_format or "")
            if v is None:
                t = ""
            elif isinstance(v, datetime.datetime):
                t = f"{v.day}/{v.month}/{v.year}" if regional == "es-AR" else f"{v.month}/{v.day}/{v.year}"
            elif isinstance(v, (int, float)) and contable:
                n = f"{abs(v):,.2f}"
                if regional == "es-AR":
                    n = n.replace(",", "_").replace(".", ",").replace("_", ".")
                t = " $ -   " if v == 0 else (f" $ {n} " if v > 0 else f"-$ {n} ")
            elif isinstance(v, str) and contable:
                t = f" {v} "
            else:
                t = str(v)
            celdas.append(t)
        lineas.append("\t".join(celdas))
    return "\r\n".join(lineas) + "\r\n"


PEGADO = portapapeles("es-AR")
if TSV_REAL:
    PEGADO_REAL = open(TSV_REAL, "rb").read().decode("utf-8")
LINEAS = PEGADO.split("\r\n")[:-1]
ENC, FILAS = LINEAS[0], LINEAS[1:]


def con_filas(*filas, encabezado=None):
    return "\r\n".join([ENC if encabezado is None else encabezado, *filas]) + "\r\n"


def cambiar(linea, columna, valor):
    celdas = linea.split("\t")
    celdas[ENCABEZADOS.index(columna)] = valor
    return "\t".join(celdas)


with db() as conn:
    for correo, tipo in (("ui.resp@test.local", "OPERADOR"), ("ui.consulta@test.local", "CONSULTOR")):
        conn.execute("DELETE FROM usuarios WHERE correo = ?", (correo,))
        conn.execute(
            "INSERT INTO usuarios (nombre_apellido, correo, password_hash, estado, tipo_usuario, modulos_json, paneles_json, acciones_json, "
            "bloqueado, intentos_fallidos) VALUES (?, ?, ?, 'ACTIVO', ?, ?, '{}', '{}', 0, 0)",
            (correo.split("@")[0], correo, hashlib.sha256(CLAVE.encode()).hexdigest(), tipo, json.dumps(["facturacion"])))
    uid = conn.execute("SELECT id FROM usuarios WHERE correo = 'ui.resp@test.local'").fetchone()[0]
    conn.execute("INSERT INTO responsables_modulo (modulo, usuario_id) VALUES ('facturacion', ?)", (uid,))
    fila_cli = conn.execute("SELECT codigo FROM clientes WHERE razon_social = ?", (RAZON,)).fetchone()
    assert fila_cli, f"El cliente {RAZON} no está en la tabla clientes"
    CLI = fila_cli["codigo"]
    OTRO = conn.execute("SELECT codigo FROM clientes WHERE codigo <> ? AND COALESCE(razon_social, '') <> '' ORDER BY codigo LIMIT 1", (CLI,)).fetchone()["codigo"]
    TABLAS_FAC = [r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table' AND name LIKE 'fac\\_%' ESCAPE '\\' ORDER BY name")]


def huella():
    """Filas de todas las tablas del módulo: pegar y preparar formularios no debe cambiarla."""
    with db() as conn:
        return {t: conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in TABLAS_FAC}


RESULTADOS, ERRORES, RESPUESTAS = [], [], []


def ok(nombre, condicion, detalle=""):
    RESULTADOS.append(f"{'OK   ' if condicion else 'FALLA'} {nombre}{f' · {detalle}' if detalle else ''}")
    return condicion


def captura(pagina, nombre):
    pagina.screenshot(path=os.path.join(CAPTURAS, f"e2e_tango_{nombre}.png"), full_page=True)


def ingresar(navegador, usuario):
    contexto = navegador.new_context(viewport={"width": 1600, "height": 1000}, locale="es-AR")
    pagina = contexto.new_page()
    pagina.on("pageerror", lambda e: ERRORES.append(f"{usuario}: {e}"))
    pagina.on("response", lambda r: RESPUESTAS.append(f"{r.status} {r.request.method} {r.url.replace(URL, '')}") if r.status >= 400 else None)
    pagina.on("dialog", lambda d: d.accept())  # confirmaciones del flujo: se aceptan como lo haría el usuario
    respuesta = contexto.request.post(f"{URL}/login", form={"username": usuario, "password": CLAVE}, max_redirects=0)
    assert respuesta.headers.get("location") == "/", f"Login fallido para {usuario}"
    pagina.goto(f"{URL}/facturacion")
    pagina.wait_for_function("typeof S !== 'undefined' && S.sesion && S.cat.clientes.length > 0")
    return pagina


def api(pagina, ruta):
    r = pagina.request.get(URL + API + ruta)
    assert r.ok, f"{ruta}: {r.status} {r.text()}"
    return r.json()


def parsear(pagina, texto, orden="auto"):
    return pagina.evaluate("([t, o]) => parsearTango(t, o)", [texto, orden])


def reconocido(res):
    return [{k: r.get(k) for k in ("tipo", "letra", "punto_venta", "numero", "fecha", "vencimiento", "debe", "haber", "saldo", "estado", "leyenda")}
            for r in res["filas"]]


def ir(pagina, vista):
    pagina.click(f".menu-btn[data-vista='{vista}']")
    pagina.wait_for_function(f"S.vista === '{vista}'")
    pagina.wait_for_timeout(300)


def abierto(pagina, modal):
    pagina.wait_for_function(f"document.getElementById('{modal}').classList.contains('open')", timeout=10000)


def cerrado(pagina, modal):
    pagina.wait_for_function(f"!document.getElementById('{modal}').classList.contains('open')", timeout=10000)


def texto(pagina, selector):
    return pagina.locator(selector).inner_text()


def fila_tango(pagina, numero):
    return pagina.locator(f"#tgo-res tr:has(td:text-is('{numero}'))")


with sync_playwright() as p:
    navegador = p.chromium.launch(channel="msedge", headless=True)
    pg = ingresar(navegador, "ui.resp@test.local")

    # ---------- 1. El Excel de GE.MAR, tal como lo pega Excel ----------
    cuenta = {t: sum(e["tipo"] == t for e in ESPERADO) for t in ("FC", "NC", "ND", "REC")}
    ok("excel: 45 movimientos (16 FAC, 11 N/C, 3 N/D, 15 REC)", len(ESPERADO) == 45 and cuenta == {"FC": 16, "NC": 11, "ND": 3, "REC": 15}, str(cuenta))
    if TSV_REAL:
        ok("portapapeles: la copia emulada es idéntica a la copia real de Excel", PEGADO == PEGADO_REAL)
        variantes = [("copia real de Excel (es-AR)", PEGADO_REAL, "dmy", ",")]
    else:
        variantes = []
    variantes += [("es-AR: d/m/aaaa y $ 1.234,56", PEGADO, "dmy", ","), ("en-US: m/d/aaaa y $ 1,234.56", portapapeles("en-US"), "mdy", ".")]
    for nombre, texto_pegado, orden, decimal in variantes:
        res = parsear(pg, texto_pegado)
        ok(f"{nombre}: 45 filas iguales al Excel (tipo, letra, PV, número, fechas, debe, haber, saldo, estado, leyenda)",
           res["ok"] and reconocido(res) == ESPERADO, res.get("error") or next((f"{a} ≠ {b}" for a, b in zip(reconocido(res), ESPERADO) if a != b), ""))
        ok(f"{nombre}: sin errores ni avisos, orden {orden}, decimal «{decimal}»",
           res["ok"] and not any(r["errores"] or r["avisos"] for r in res["filas"]) and res["orden"] == orden and res["decimal"] == decimal and res["razon_social"] == RAZON)
        ok(f"{nombre}: saldo anterior 0, final 80.569.304,36 y al {CORTE} 15.807.729,91",
           res["ok"] and res["saldo_anterior"] == 0 and res["saldo_final"] == 80569304.36
           and pg.evaluate("([t, d]) => tangoSaldoAl(parsearTango(t), d)", [texto_pegado, CORTE]) == 15807729.91)
        ok(f"{nombre}: los importes salen de su columna (FAC/N/D en DEBE; N/C/REC en HABER)",
           res["ok"] and all(r["importe"] == (e["debe"] if e["tipo"] in ("FC", "ND") else e["haber"]) for r, e in zip(res["filas"], ESPERADO)))
    # Mismas filas con columnas en otro orden, una columna extra, fechas ISO y saltos de línea LF.
    orden_nuevo = ["N_COMP", "T_COMP", "SALDO", "HABER", "DEBE", "FECHA_VTO", "FECHA_EMIS", "EXTRA", "LEYENDA", "ESTADO", "RAZON_SOC"]
    filas_iso = []
    for e, linea in zip(ESPERADO, FILAS):
        c = dict(zip(ENCABEZADOS, linea.split("\t")))
        c.update({"FECHA_EMIS": e["fecha"], "FECHA_VTO": e["vencimiento"] or "", "EXTRA": "x"})
        filas_iso.append("\t".join(c[k] for k in orden_nuevo))
    res = parsear(pg, "\n".join(["\t".join(orden_nuevo), *filas_iso]))
    ok("columnas en otro orden, columna extra, fechas ISO y LF: mismo resultado", res["ok"] and reconocido(res) == ESPERADO, res.get("error", ""))

    # ---------- 2. Formatos incorrectos: el pegado entero no se reconoce ----------
    sin_tab = PEGADO.replace("\t", ";")
    casos = [
        ("pegado vacío", "", "No hay nada pegado"),
        ("texto sin columnas (PDF o CSV)", sin_tab, "No se reconocen columnas"),
        ("filas sin la fila de encabezados", "\r\n".join(FILAS[:3]), "faltan: FECHA_EMIS"),
        ("encabezados de otro reporte (FECHA, TIPO, NUMERO)", "FECHA\tTIPO\tNUMERO\tDEBE\tHABER\n2/7/2025\tFAC\tA0000600000474\t1\t0", "faltan: FECHA_EMIS, T_COMP, N_COMP, SALDO"),
        ("falta la columna SALDO", con_filas(*["\t".join(x.split("\t")[:8]) for x in FILAS[:2]], encabezado="\t".join(ENC.split("\t")[:8])), "faltan: SALDO"),
        ("encabezado repetido", con_filas(FILAS[0] + "\t1", encabezado=ENC + "\tDEBE"), "Encabezado repetido: DEBE"),
        ("solo la fila de encabezados", ENC, "Solo se pegó la fila de encabezados"),
        ("dos clientes en el mismo pegado", con_filas(FILAS[4], cambiar(FILAS[5], "RAZON_SOC", "OTRA S.A.")), "más de un cliente"),
        ("decimales mezclados (coma y punto)", con_filas(FILAS[4], cambiar(cambiar(FILAS[5], "HABER", " $ 24,248,925.95 "), "SALDO", " $ 145,686,741.64 ")), "mezclan coma y punto"),
        ("fechas mezcladas (día/mes y mes/día)", con_filas(FILAS[5], cambiar(FILAS[6], "FECHA_EMIS", "8/14/2025")), "mezclan día/mes y mes/día"),
    ]
    for nombre, texto_pegado, mensaje in casos:
        res = parsear(pg, texto_pegado)
        ok(f"formato incorrecto: {nombre}", not res["ok"] and mensaje in res["error"] and not res["filas"], res.get("error", "reconocido"))
    ambiguo = con_filas(*FILAS[:4])  # 2/7, 3/7 y 11/7: no se sabe si es día/mes o mes/día
    res = parsear(pg, ambiguo)
    ok("fechas ambiguas: pide elegir el formato", not res["ok"] and res.get("ambiguo") and "Elegí el formato" in res["error"], res.get("error", ""))
    ok("fechas ambiguas: día/mes → 02/07/2025, mes/día → 07/02/2025",
       parsear(pg, ambiguo, "dmy")["filas"][0]["fecha"] == "2025-07-02" and parsear(pg, ambiguo, "mdy")["filas"][0]["fecha"] == "2025-02-07")

    # ---------- 3. Filas con errores: no se preparan; el resto sí ----------
    errores_fila = [
        ("tipo no reconocido", cambiar(FILAS[0], "T_COMP", "FCE"), "tipo «FCE» no reconocido"),
        ("número mal formado", cambiar(FILAS[0], "N_COMP", "A06-474"), "número «A06-474» no reconocido"),
        ("factura sin letra", cambiar(FILAS[0], "N_COMP", "0000600000474"), "falta la letra"),
        ("letra inválida", cambiar(FILAS[0], "N_COMP", "X0000600000474"), "letra «X» no válida"),
        ("recibo con letra", cambiar(FILAS[3], "N_COMP", "A0000100000446"), "un recibo no lleva letra"),
        ("fecha inexistente", cambiar(FILAS[0], "FECHA_EMIS", "31/2/2025"), "fecha de emisión «31/2/2025» no reconocida"),
        ("vencimiento no reconocido", cambiar(FILAS[0], "FECHA_VTO", "mañana"), "vencimiento «mañana» no reconocido"),
        ("importe no numérico", cambiar(FILAS[0], "DEBE", " $ abc "), "DEBE «$ abc» no es un importe"),
        ("importe con 3 decimales", cambiar(FILAS[0], "SALDO", " $ 530.385.727,401 "), "SALDO «$ 530.385.727,401» no es un importe"),
        ("factura con el importe en HABER", cambiar(cambiar(FILAS[0], "DEBE", " $ -   "), "HABER", " $ 530.385.727,40 "), "FAC lleva su importe solo en DEBE"),
        ("recibo con el importe en DEBE", cambiar(cambiar(FILAS[3], "HABER", " $ -   "), "DEBE", " $ 533.981.562,84 "), "REC lleva su importe solo en HABER"),
        ("factura en cero", cambiar(FILAS[0], "DEBE", " $ -   "), "FAC lleva su importe solo en DEBE"),
    ]
    for nombre, linea, mensaje in errores_fila:
        res = parsear(pg, con_filas(linea, *FILAS[1:6]))
        ok(f"fila con error: {nombre}", res["ok"] and any(mensaje in x for x in res["filas"][0]["errores"])
           and not any(r["errores"] for r in res["filas"][1:]), json.dumps(res["filas"][0]["errores"] if res["ok"] else res["error"], ensure_ascii=False))
    res = parsear(pg, con_filas(FILAS[0], FILAS[1], FILAS[3], FILAS[4]))  # falta la N/C de la línea 4
    ok("fila faltante: aviso de SALDO que no sigue al anterior", any("el SALDO no sigue" in a for a in res["filas"][2]["avisos"]), json.dumps(res["filas"][2]["avisos"], ensure_ascii=False))
    # Pegados cortos con fechas hasta el 12: el formato se indica, como lo haría el usuario en pantalla.
    res = parsear(pg, con_filas(FILAS[0], FILAS[1], FILAS[1]), "dmy")
    ok("número repetido: aviso en las dos filas", res["filas"][1]["repetido"] == 2 and res["filas"][2]["repetido"] == 2)
    res = parsear(pg, con_filas(FILAS[1], FILAS[0]), "dmy")
    ok("fila fuera de orden de fecha o saldo: aviso", any("SALDO no sigue" in a for a in res["filas"][1]["avisos"]))
    total = "\t" * 6 + " $ 1 \t $ 1 \t $ 1 \t\t"
    res = parsear(pg, con_filas(*FILAS[:3], total), "dmy")
    ok("fila de totales sin tipo ni número: se ignora", res["ok"] and len(res["filas"]) == 3 and res["ignoradas"] == [5], str(res.get("ignoradas")))
    res = parsear(pg, con_filas(cambiar(FILAS[3], "LEYENDA", '"FACT. 6-475\nY 6-474"')), "dmy")
    ok("celda con salto de línea entre comillas (como la copia Excel)", res["ok"] and res["filas"][0]["leyenda"] == "FACT. 6-475\nY 6-474")

    # ---------- 4. Pantalla: pegar no guarda nada ----------
    corte0 = api(pg, "/corte")["corte"]
    if corte0:  # la base real ya tiene el corte en borrador sin fecha: se reutiliza
        ok("corte de origen en borrador sin fecha", corte0["estado"] == "BORRADOR" and corte0["fecha_corte"] is None, json.dumps(corte0))
    else:
        r = pg.request.post(f"{URL}{API}/corte", data={})
        assert r.ok, r.text()
    r = pg.request.put(f"{URL}{API}/corte/fecha", data={"fecha_corte": CORTE})
    assert r.ok, r.text()
    pg.reload()
    pg.wait_for_function("typeof S !== 'undefined' && S.sesion && S.sesion.corte && S.cat.clientes.length > 0")
    ir(pg, "saldos")
    pg.wait_for_selector("#pm-clientes tr.click")
    ok("puesta en marcha: botón «Pegar resumen de Tango»", pg.locator("#corte-acciones button:has-text('Pegar resumen de Tango')").count() == 1)
    pg.fill("#pm-q", CLI)
    pg.click(f"#pm-clientes tr.click:has-text('{CLI}')")
    abierto(pg, "m-ini")
    pg.wait_for_function("document.getElementById('ini-pie').innerText.includes('Pegar resumen de Tango')")
    base = huella()
    pg.click("#ini-pie button:has-text('Pegar resumen de Tango')")
    abierto(pg, "m-tango")
    ok("desde la ficha: cliente elegido y moneda sin confirmar", pg.input_value("#tgo-cli") == CLI and pg.input_value("#tgo-mon") == "")
    pg.fill("#tgo-txt", PEGADO)
    pg.wait_for_selector("#tgo-res table")
    resumen = texto(pg, "#tgo-res")
    ok("pantalla: 45 filas (16 FAC · 11 N/C · 3 N/D · 15 REC), fechas día/mes/año", "45 fila(s) reconocida(s): 16 FAC · 11 N/C · 3 N/D · 15 REC" in resumen and "día/mes/año" in resumen, resumen[:300])
    ok("pantalla: saldo al inicio operativo 15.807.729,91 según el SALDO de Tango", "saldo al inicio operativo (28/02/2026): 15.807.729,91" in resumen)
    ok("sin moneda confirmada: ningún formulario se puede preparar",
       pg.locator("#tgo-res button").count() > 0 and pg.locator("#tgo-res button:enabled").count() == 0 and "Confirmá la moneda" in resumen)
    pg.select_option("#tgo-mon", "ARS")
    pg.wait_for_function("document.querySelectorAll('#tgo-res button:enabled').length > 0")
    ok("moneda confirmada: 45 filas preparables + saldo según Tango", pg.locator("#tgo-res button:enabled").count() == 46, str(pg.locator("#tgo-res button:enabled").count()))
    ok("hasta el inicio: saldo inicial; después: operación nueva",
       "Saldo inicial" in fila_tango(pg, "A-0006-00000504").inner_text() and "Operación nueva" in fila_tango(pg, "A-0006-00000508").inner_text()
       and "Anticipo…" in fila_tango(pg, "0001-00000477").inner_text() and "Cobranza…" in fila_tango(pg, "0001-00000487").inner_text())
    captura(pg, "01_pegado")
    ok("pegar y confirmar la moneda no guardó nada", huella() == base)

    # ---------- 5. Factura abierta al inicio: formulario preparado, se confirma a mano ----------
    fila_tango(pg, "A-0006-00000504").locator("button").click()
    abierto(pg, "m-doc")
    pg.wait_for_function("DOC_MODO === 'SALDO_INICIAL' && document.getElementById('d-nro').value === '504' && !document.getElementById('d-tango').classList.contains('hidden')")
    campos = pg.evaluate("""() => ({ tipo: TIPO_DOC, letra: $('d-letra').value, pv: $('d-pv').value, nro: $('d-nro').value, cli: $('d-cli').value,
      mon: $('d-mon').value, fe: $('d-fe').value, vto: $('d-vto').value, pdf: $('d-pdf').value, netos: DL.map(l => l.neto), pagos: PP.length,
      nota: $('d-si-nota').value, vinculos: VL.length, aviso: !$('d-tango').classList.contains('hidden') })""")
    ok("formulario: tipo, número, cliente, moneda confirmada, fechas y total de Tango en el control",
       campos["tipo"] == "FC" and campos["letra"] == "A" and campos["pv"] == "6" and campos["nro"] == "504" and campos["cli"] == CLI and campos["mon"] == "ARS"
       and campos["fe"] == "2026-02-07" and campos["vto"] == "2026-02-08" and campos["pdf"] == "573.714.967,50" and campos["aviso"], json.dumps(campos))
    ok("formulario: sin IVA, pagos previos ni vínculos deducidos", campos["netos"] == [""] and campos["pagos"] == 0 and campos["vinculos"] == 0, json.dumps(campos))
    recibos_al_inicio = sum(1 for e in ESPERADO if e["tipo"] == "REC" and e["fecha"] <= CORTE)
    ok(f"formulario: los {recibos_al_inicio} REC hasta el inicio, como ayuda para pagos previos", pg.locator("#d-tango a").count() == recibos_al_inicio)
    captura(pg, "02_factura_preparada")
    pg.click("#m-doc .modal-f button:has-text('Cancelar')")
    cerrado(pg, "m-doc")
    ok("preparar y cancelar no guardó nada", huella() == base)

    fila_tango(pg, "A-0006-00000504").locator("button").click()
    pg.wait_for_function("document.getElementById('m-doc').classList.contains('open') && document.getElementById('d-nro').value === '504' && !document.getElementById('d-tango').classList.contains('hidden')")
    i481 = pg.evaluate("TANGO.res.filas.findIndex(x => x.tipo === 'REC' && x.numero === 481)")
    pg.click(f"#d-tango a[onclick='tangoPagoPrevio({i481})']")
    pp = pg.evaluate("PP[0]")
    ok("REC 481 como pago previo: fecha, recibo y total de Tango; importe aplicado vacío",
       pp == {"fecha_pago": "2026-02-27", "moneda_pago": "ARS", "importe_pago": "", "tipo_cambio": "", "recibo_punto_venta": "1", "recibo_numero": "481",
              "recibo_total": "400.000.000,00"}, json.dumps(pp))
    pg.click("#d-ok")  # sin neto: el formulario no se registra
    pg.wait_for_function("document.getElementById('d-val').innerText.includes('No se registró')")
    rechazos = [r for r in RESPUESTAS if r == f"400 POST {API}/documentos/validar"]
    RESPUESTAS[:] = [r for r in RESPUESTAS if r not in rechazos]  # rechazo esperado: se controla acá
    ok("sin completar el IVA no se registra (validación de siempre)", huella() == base and len(rechazos) == 1, texto(pg, "#d-val")[:160])
    pg.fill("#d-lineas tr:nth-child(1) td:nth-child(3) input", "474144601,24")
    pg.fill("#d-si-pagos tr:nth-child(1) td:nth-child(3) input", "400000000")
    pg.click("#d-si-pagos-w button:has-text('+ Pago previo')")
    pg.fill("#d-si-pagos tr:nth-child(2) td:nth-child(1) input", "2026-02-07")
    pg.fill("#d-si-pagos tr:nth-child(2) td:nth-child(3) input", "157907237,59")
    ok("saldo al inicio en pantalla = SALDO de Tango", "saldoalinicioARS15.807.729,91" in texto(pg, "#d-si-res").replace(" ", ""), texto(pg, "#d-si-res"))
    pg.click("#d-ok")
    cerrado(pg, "m-doc")
    doc = next(d for d in api(pg, f"/documentos?cliente={CLI}") if d["numero"] == 504)
    ok("factura registrada al confirmar: total y saldo según Tango", doc["total"] == 573714967.5 and doc["saldo"] == 15807729.91 and doc["origen"] == "SALDO_INICIAL",
       f"total {doc['total']} saldo {doc['saldo']}")
    pg.wait_for_function("document.querySelector('#tgo-res') && document.querySelector('#tgo-res').innerText.includes('YA REGISTRADO')")
    ok("la fila queda marcada como ya registrada", "YA REGISTRADO" in fila_tango(pg, "A-0006-00000504").inner_text() and pg.locator("#tgo-res .badge:has-text('YA REGISTRADO')").count() == 1)

    # ---------- 6. Saldo según Tango al inicio: se confirma en su formulario ----------
    pg.click("#tgo-res button:has-text('como saldo según Tango al inicio')")
    abierto(pg, "m-conc")
    ok("conciliación preparada con el SALDO de Tango al inicio", pg.input_value("#k-cli") == CLI and pg.input_value("#k-mon") == "ARS" and pg.input_value("#k-saldo") == "15.807.729,91")
    antes = huella()
    pg.click("#m-conc .modal-f button:has-text('Guardar')")
    cerrado(pg, "m-conc")
    saldos = api(pg, f"/puesta-en-marcha/{CLI}")["saldos"]
    ok("conciliado con Tango al confirmar", any(s["moneda"] == "ARS" and s["conciliada"] and s["saldo_tango"] == 15807729.91 for s in saldos) and huella() != antes, json.dumps(saldos))

    # ---------- 7. Otros clientes: aviso de RAZON_SOC distinta ----------
    pg.select_option("#tgo-cli", OTRO)
    pg.wait_for_function("document.getElementById('tgo-res').innerText.includes('no coincide')")
    ok("otro cliente: aviso y formularios bloqueados", pg.locator("#tgo-res tr button:enabled").count() == 0)
    pg.click("#tgo-conf")
    pg.wait_for_function("TANGO.confirma && document.querySelectorAll('#tgo-res tr button:enabled').length > 0")
    ok("otro cliente: se habilita solo al confirmarlo", pg.locator("#tgo-res tr button:enabled").count() > 0)
    pg.select_option("#tgo-cli", CLI)
    pg.wait_for_function("!document.getElementById('tgo-res').innerText.includes('no coincide') && document.querySelectorAll('#tgo-res tr button:enabled').length > 0")

    # ---------- 8. Cliente validado: operaciones nuevas preparadas como formularios normales ----------
    pg.click("#m-tango .modal-f button:has-text('Cerrar')")
    cerrado(pg, "m-tango")
    pg.click("#ini-pie button:has-text('Validar con antecedentes')")
    pg.wait_for_function("document.getElementById('ini-tit').innerText.includes('HABILITADO')")
    pg.click("#ini-pie button:has-text('Pegar resumen de Tango')")
    abierto(pg, "m-tango")
    pg.wait_for_selector("#tgo-res table")
    ok("al reabrir: el texto sigue pegado y la moneda se vuelve a confirmar", pg.input_value("#tgo-mon") == "" and pg.locator("#tgo-res button:enabled").count() == 0)
    pg.select_option("#tgo-mon", "ARS")
    pg.wait_for_function("document.querySelectorAll('#tgo-res button:enabled').length > 0")
    base = huella()
    fila_tango(pg, "A-0006-00000508").locator("button").click()
    pg.wait_for_function("document.getElementById('m-doc').classList.contains('open') && DOC_MODO === 'NORMAL' && document.getElementById('d-nro').value === '508' && !document.getElementById('d-tango').classList.contains('hidden')")
    ok("operación nueva: formulario normal, sin bloque de saldo inicial",
       pg.evaluate("$('d-si-w').classList.contains('hidden') && $('d-fe').value === '2026-04-04' && $('d-vto').value === '2026-04-11' && $('d-pdf').value === '46.081.736,80'"))
    pg.click("#m-doc .modal-f button:has-text('Cancelar')")
    cerrado(pg, "m-doc")
    fila_tango(pg, "0001-00000487").locator("button").click()
    abierto(pg, "m-cob")
    pg.wait_for_function("document.getElementById('c-nro').value === '487' && !document.getElementById('c-tango').classList.contains('hidden')")
    cob = pg.evaluate("""() => ({ modo: COB_MODO, pv: $('c-pv').value, fecha: $('c-fecha').value, mon: $('c-mon').value, cli: $('c-cli').value,
      valores: VALORES.map(v => v.importe), asignado: [...document.querySelectorAll('#c-imp input.imp')].filter(i => i.value).length,
      pendientes: document.querySelectorAll('#c-imp input.imp').length, aviso: $('c-tango').innerText })""")
    ok("cobranza preparada: recibo, fecha, moneda confirmada y total de Tango",
       cob["modo"] == "NORMAL" and cob["pv"] == "1" and cob["fecha"] == "2026-04-08" and cob["mon"] == "ARS" and cob["cli"] == CLI and cob["valores"] == ["491.736.880,60"], json.dumps(cob, ensure_ascii=False))
    ok("cobranza preparada: sin aplicación a facturas deducida (la leyenda solo se muestra)",
       cob["asignado"] == 0 and cob["pendientes"] >= 1 and "FACT. 5-203" in cob["aviso"], json.dumps(cob, ensure_ascii=False))
    captura(pg, "03_cobranza_preparada")
    pg.click("#m-cob .modal-f button:has-text('Cancelar')")
    cerrado(pg, "m-cob")
    ok("preparar operaciones nuevas y cancelar no guardó nada", huella() == base)

    # ---------- 9. Formato incorrecto y fechas ambiguas en pantalla ----------
    pg.fill("#tgo-txt", sin_tab)
    pg.wait_for_function("document.getElementById('tgo-res').innerText.includes('No se reconocieron las filas')")
    ok("pantalla: pegado sin columnas → error y ninguna fila", pg.locator("#tgo-res table").count() == 0)
    pg.fill("#tgo-txt", ambiguo)
    pg.wait_for_function("document.getElementById('tgo-res').innerText.includes('Elegí el formato')")
    pg.select_option("#tgo-ord", "dmy")
    pg.wait_for_selector("#tgo-res table")
    ok("pantalla: fechas ambiguas → se reconocen al elegir día/mes/año", "4 fila(s) reconocida(s)" in texto(pg, "#tgo-res") and "02/07/2025" in texto(pg, "#tgo-res"))
    pg.select_option("#tgo-ord", "auto")
    pg.fill("#tgo-txt", con_filas(cambiar(FILAS[0], "T_COMP", "FCE"), FILAS[1], FILAS[4]))  # 14/8: fechas día/mes
    pg.wait_for_function("document.getElementById('tgo-res').innerText.includes('con errores')")
    ok("pantalla: fila con error sin botón; las demás preparables",
       fila_tango(pg, "2").locator("button").count() == 0 and "tipo «FCE» no reconocido" in texto(pg, "#tgo-res") and pg.locator("#tgo-res tr button:enabled").count() == 2)
    pg.click("#m-tango .modal-f button:has-text('Cerrar')")

    # ---------- 10. Perfil de consulta: sin pegado ----------
    cons = ingresar(navegador, "ui.consulta@test.local")
    ir(cons, "saldos")
    cons.wait_for_selector("#pm-clientes tr.click")
    cons.fill("#pm-q", CLI)
    cons.click(f"#pm-clientes tr.click:has-text('{CLI}')")
    abierto(cons, "m-ini")
    ok("consulta: sin botón «Pegar resumen de Tango»", cons.locator("button:has-text('Pegar resumen de Tango')").count() == 0)
    navegador.close()

servidor.should_exit = True
ok("la base de origen no cambió", hashlib.sha256(open(DB_ORIGEN, "rb").read()).hexdigest() == SHA_ORIGEN)
ok("el Excel no cambió", hashlib.sha256(open(XLSX, "rb").read()).hexdigest() == SHA_XLSX)
for linea in RESULTADOS:
    print(linea)
print("Base temporal:", DB)
print("Capturas:", CAPTURAS)
print("EXCEPCIONES JS:", json.dumps(ERRORES, ensure_ascii=False))
print("Respuestas con error:", json.dumps(RESPUESTAS, ensure_ascii=False))
sys.exit(0 if all(r.startswith("OK") for r in RESULTADOS) and not ERRORES and not RESPUESTAS else 1)
